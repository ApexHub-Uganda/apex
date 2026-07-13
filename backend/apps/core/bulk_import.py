"""Bulk import utilities — Excel and CSV template generation, parsing, and validation."""
from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any, Callable

from apps.core.email_validation import validate_deliverable_email

EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
PHONE_RE = re.compile(r"^\+?[\d\s\-()]{7,20}$")

SUPPORTED_IMPORT_EXTENSIONS = (".xlsx", ".xlsm", ".csv", ".txt")
DEFAULT_IMPORT_FORMAT = "xlsx"
EXCEL_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CSV_CONTENT_TYPE = "text/csv; charset=utf-8"


@dataclass
class ImportColumn:
    key: str
    label: str
    required: bool = False
    help_text: str = ""
    choices: list[str] | None = None
    field_type: str = "string"  # string, date, decimal, integer, boolean, email, phone


@dataclass
class ImportSpec:
    entity_name: str
    columns: list[ImportColumn]
    description: str = ""

    @property
    def required_keys(self) -> set[str]:
        return {c.key for c in self.columns if c.required}

    @property
    def all_keys(self) -> set[str]:
        return {c.key for c in self.columns}

    def column_map(self) -> dict[str, ImportColumn]:
        return {c.key: c for c in self.columns}


def normalize_header(header: str) -> str:
    return header.strip().lower().replace(" ", "_").replace("-", "_")


def _hint_row(spec: ImportSpec) -> list[str]:
    hints = []
    for col in spec.columns:
        hint = col.help_text or ""
        if col.choices:
            hint = f"{hint} Options: {', '.join(col.choices)}".strip()
        if col.required:
            hint = f"{hint} (required)".strip()
        hints.append(hint)
    return hints


def is_supported_import_filename(filename: str) -> bool:
    lowered = (filename or "").lower()
    return any(lowered.endswith(ext) for ext in SUPPORTED_IMPORT_EXTENSIONS)


def unsupported_import_message() -> str:
    return "Upload an Excel workbook (.xlsx) or CSV file (.csv)."


def generate_csv_template(spec: ImportSpec) -> bytes:
    """Build a UTF-8 CSV template with header row and hint row."""
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow([col.label for col in spec.columns])
    writer.writerow(_hint_row(spec))
    return output.getvalue().encode("utf-8-sig")


def generate_excel_template(spec: ImportSpec) -> bytes:
    """Build an Excel workbook template with header and hint rows."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Import"
    worksheet.append([col.label for col in spec.columns])
    worksheet.append(_hint_row(spec))
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def generate_import_template(spec: ImportSpec, *, file_format: str = DEFAULT_IMPORT_FORMAT) -> tuple[bytes, str, str]:
    """Return template bytes, filename, and content type."""
    slug = spec.entity_name.lower().replace(" ", "_")
    if file_format == "csv":
        return generate_csv_template(spec), f"{slug}_import_template.csv", CSV_CONTENT_TYPE
    return generate_excel_template(spec), f"{slug}_import_template.xlsx", EXCEL_CONTENT_TYPE


def _cell_to_str(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def _is_hint_row(row_values: list[str]) -> bool:
    if not row_values:
        return False
    return all(
        not value
        or value.startswith("(")
        or "required" in value.lower()
        or "options:" in value.lower()
        for value in row_values
        if value
    )


def _rows_to_records(raw_rows: list[list[str]]) -> tuple[list[str], list[dict[str, str]]]:
    if not raw_rows:
        return [], []

    raw_headers = raw_rows[0]
    headers = [normalize_header(h) for h in raw_headers]
    data_rows: list[dict[str, str]] = []
    for row_values in raw_rows[1:]:
        normalized_values = [value.strip() for value in row_values]
        if not any(normalized_values):
            continue
        if _is_hint_row(normalized_values):
            continue
        row_dict: dict[str, str] = {}
        for idx, header in enumerate(headers):
            if not header:
                continue
            row_dict[header] = normalized_values[idx] if idx < len(normalized_values) else ""
        data_rows.append(row_dict)
    return headers, data_rows


def _parse_csv_bytes(raw: bytes) -> tuple[list[str], list[dict[str, str]]]:
    text = raw.decode("utf-8-sig", errors="replace")
    reader = csv.reader(io.StringIO(text))
    rows = [[cell.strip() for cell in row] for row in reader]
    return _rows_to_records(rows)


def _parse_excel_bytes(raw: bytes) -> tuple[list[str], list[dict[str, str]]]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    worksheet = workbook.active
    rows = [
        [_cell_to_str(cell) for cell in row]
        for row in worksheet.iter_rows(values_only=True)
    ]
    workbook.close()
    return _rows_to_records(rows)


def _detect_upload_format(filename: str, raw: bytes) -> str:
    lowered = (filename or "").lower()
    if lowered.endswith((".xlsx", ".xlsm")):
        return "excel"
    if lowered.endswith((".csv", ".txt")):
        return "csv"
    if raw[:2] == b"PK":
        return "excel"
    return "csv"


def parse_upload(file_obj) -> tuple[list[str], list[dict[str, str]]]:
    """Parse uploaded Excel (.xlsx) or CSV into normalized row dicts."""
    raw = file_obj.read()
    if not isinstance(raw, bytes):
        raw = str(raw).encode("utf-8-sig", errors="replace")

    filename = getattr(file_obj, "name", "") or ""
    if _detect_upload_format(filename, raw) == "excel":
        return _parse_excel_bytes(raw)
    return _parse_csv_bytes(raw)


def map_headers_to_keys(headers: list[str], spec: ImportSpec) -> dict[str, str]:
    """Map file headers to spec column keys (by label or key)."""
    col_map = spec.column_map()
    label_lookup = {normalize_header(c.label): c.key for c in spec.columns}
    key_lookup = {normalize_header(c.key): c.key for c in spec.columns}
    mapping: dict[str, str] = {}
    for header in headers:
        if header in key_lookup:
            mapping[header] = key_lookup[header]
        elif header in label_lookup:
            mapping[header] = label_lookup[header]
    return mapping


def validate_headers(headers: list[str], spec: ImportSpec) -> list[dict[str, Any]]:
    """Return header-level errors (missing required columns)."""
    mapping = map_headers_to_keys(headers, spec)
    mapped_keys = set(mapping.values())
    errors: list[dict[str, Any]] = []
    for col in spec.columns:
        if col.required and col.key not in mapped_keys:
            errors.append({
                "row": 0,
                "field": col.key,
                "message": f"Missing required column: {col.label}",
            })
    unknown = [h for h in headers if h and h not in mapping]
    if unknown:
        for h in unknown:
            errors.append({
                "row": 0,
                "field": h,
                "message": f"Unrecognized column: {h}",
                "severity": "warning",
            })
    return errors


def _parse_bool(value: str) -> bool | None:
    v = value.strip().lower()
    if v in ("yes", "y", "true", "1"):
        return True
    if v in ("no", "n", "false", "0"):
        return False
    return None


def _parse_date(value: str) -> date | None:
    v = value.strip()
    if not v:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(v, fmt).date()
        except ValueError:
            continue
    return None


def _coerce_value(col: ImportColumn, raw: str) -> tuple[Any, str | None]:
    value = raw.strip()
    if not value:
        if col.required:
            return None, f"{col.label} is required"
        return None, None

    if col.key == "gender":
        normalized = value.strip().upper()
        if normalized in ("M", "MALE"):
            return "male", None
        if normalized in ("F", "FEMALE"):
            return "female", None
        return None, f"{col.label} must be M or F"

    if col.choices:
        normalized = value.strip().upper().replace(" ", "_")
        valid = {c.upper(): c for c in col.choices}
        if normalized not in valid:
            return None, f"{col.label} must be one of: {', '.join(col.choices)}"
        return valid[normalized], None

    if col.field_type == "date":
        parsed = _parse_date(value)
        if parsed is None:
            return None, f"{col.label} must be a valid date (YYYY-MM-DD)"
        return parsed, None

    if col.field_type == "decimal":
        try:
            return Decimal(value.replace(",", "")), None
        except InvalidOperation:
            return None, f"{col.label} must be a valid number"

    if col.field_type == "integer":
        try:
            return int(value), None
        except ValueError:
            return None, f"{col.label} must be a whole number"

    if col.field_type == "boolean":
        parsed = _parse_bool(value)
        if parsed is None:
            return None, f"{col.label} must be yes/no"
        return parsed, None

    if col.field_type == "email":
        error = validate_deliverable_email(value, required=col.required)
        if error:
            return None, error
        return value.lower(), None

    if col.field_type == "phone":
        if not PHONE_RE.match(value):
            return None, f"{col.label} must be a valid phone number"
        return value, None

    return value, None


def validate_rows(
    spec: ImportSpec,
    headers: list[str],
    raw_rows: list[dict[str, str]],
    *,
    row_resolver: Callable[[dict[str, Any], int], list[dict[str, Any]]] | None = None,
) -> dict[str, Any]:
    """Validate parsed rows against spec; optional resolver adds FK/business rules."""
    header_errors = validate_headers(headers, spec)
    mapping = map_headers_to_keys(headers, spec)
    col_map = spec.column_map()

    ready: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = list(header_errors)

    if any(e.get("severity") != "warning" for e in header_errors):
        return {
            "ready": [],
            "errors": errors,
            "total_rows": len(raw_rows),
            "valid_count": 0,
            "error_count": len(errors),
            "can_commit": False,
        }

    for idx, raw_row in enumerate(raw_rows, start=2):
        parsed: dict[str, Any] = {"_row_number": idx}
        row_errors: list[dict[str, Any]] = []

        for file_header, spec_key in mapping.items():
            col = col_map.get(spec_key)
            if not col:
                continue
            raw_val = raw_row.get(file_header, "")
            coerced, err = _coerce_value(col, raw_val)
            if err:
                row_errors.append({"row": idx, "field": spec_key, "message": err})
            elif coerced is not None:
                parsed[spec_key] = coerced

        for col in spec.columns:
            if col.required and col.key not in parsed:
                row_errors.append({
                    "row": idx,
                    "field": col.key,
                    "message": f"{col.label} is required",
                })

        if row_resolver:
            row_errors.extend(row_resolver(parsed, idx))

        if row_errors:
            errors.extend(row_errors)
        else:
            ready.append(parsed)

    return {
        "ready": ready,
        "errors": errors,
        "total_rows": len(raw_rows),
        "valid_count": len(ready),
        "error_count": len([e for e in errors if e.get("severity") != "warning"]),
        "can_commit": len(ready) > 0 and not any(
            e.get("severity") != "warning" and e.get("row") == 0 for e in errors
        ),
    }