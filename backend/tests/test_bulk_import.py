"""Tests for CSV bulk import infrastructure."""
from __future__ import annotations

import io

import pytest
from rest_framework.test import APIClient

from apps.academics.models import AcademicYear, Class, Stream
from apps.core.bulk_import import (
    EXCEL_CONTENT_TYPE,
    generate_csv_template,
    generate_excel_template,
    generate_import_template,
    is_supported_import_filename,
    parse_upload,
    validate_rows,
)
from apps.core.constants import UserRole
from apps.staff.import_handlers import STAFF_IMPORT_SPEC
from apps.staff.services import onboard_staff
from apps.students.import_handlers import (
    PARENT_IMPORT_SPEC,
    STUDENT_IMPORT_SPEC,
    StudentImportContext,
    commit_parent_rows,
    commit_student_rows,
    get_student_import_context,
    student_import_resolver,
)
from apps.students.models import Parent, Student
from apps.students.profile import IMPORT_PLACEHOLDER_DOB, is_profile_incomplete
from apps.subscriptions.services import assign_plan_features

VALID_STUDENT_EMAIL = "student.import@test.edu"
VALID_STUDENT_PHONE = "0712345678"


@pytest.fixture
def import_plan(plan):
    assign_plan_features(plan, [
        "student_management", "parent_management", "staff_management",
        "classes", "academic_years", "terms", "student_billing",
        "hr_departments", "leave_requests",
    ])
    return plan


@pytest.fixture
def academic_setup(db, tenant):
    year = AcademicYear.objects.create(
        tenant=tenant, name="2026", start_date="2026-01-01", end_date="2026-12-31", is_current=True,
    )
    school_class = Class.objects.create(
        tenant=tenant, name="Grade 7", code="G7A", academic_year=year, capacity=40,
    )
    stream = Stream.objects.create(tenant=tenant, school_class=school_class, name="East")
    return year, school_class, stream


def test_is_supported_import_filename():
    assert is_supported_import_filename("students.xlsx") is True
    assert is_supported_import_filename("students.xlsm") is True
    assert is_supported_import_filename("students.csv") is True
    assert is_supported_import_filename("students.txt") is True
    assert is_supported_import_filename("students.pdf") is False


def test_generate_excel_template_has_minimal_headers():
    content = generate_excel_template(STUDENT_IMPORT_SPEC)
    assert content[:2] == b"PK"
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(content), read_only=True)
    rows = list(workbook.active.iter_rows(values_only=True))
    workbook.close()
    assert rows[0] == ("First Name", "Last Name", "Sex")


def test_generate_import_template_defaults_to_excel():
    content, filename, content_type = generate_import_template(STUDENT_IMPORT_SPEC)
    assert filename.endswith(".xlsx")
    assert content_type == EXCEL_CONTENT_TYPE
    assert content[:2] == b"PK"


def test_generate_csv_template_has_minimal_headers():
    content = generate_csv_template(STUDENT_IMPORT_SPEC).decode("utf-8-sig")
    assert "First Name" in content
    assert "Last Name" in content
    assert "Email" not in content.splitlines()[0]
    assert "Phone" not in content.splitlines()[0]
    assert "Admission Number" not in content
    assert "Class Code" not in content


def test_parent_template_is_minimal():
    content = generate_csv_template(PARENT_IMPORT_SPEC).decode("utf-8-sig")
    assert "M-Pesa" not in content
    assert "First Name" in content
    assert "Email" in content
    assert "Phone" in content


def test_staff_template_has_contact_columns():
    content = generate_csv_template(STAFF_IMPORT_SPEC).decode("utf-8-sig")
    header = content.splitlines()[0]
    assert "First Name" in header
    assert "Last Name" in header
    assert "Email" in header
    assert "Phone" in header
    assert "Date Joined" not in header


def test_import_templates_differ_by_entity():
    student_header = generate_csv_template(STUDENT_IMPORT_SPEC).decode("utf-8-sig").splitlines()[0]
    parent_header = generate_csv_template(PARENT_IMPORT_SPEC).decode("utf-8-sig").splitlines()[0]
    staff_header = generate_csv_template(STAFF_IMPORT_SPEC).decode("utf-8-sig").splitlines()[0]
    assert student_header != parent_header
    assert student_header != staff_header
    assert parent_header == staff_header


def test_validate_headers_missing_required():
    headers, rows = parse_upload(io.BytesIO(b"first_name\nJohn\n"))
    result = validate_rows(STUDENT_IMPORT_SPEC, headers, rows)
    assert result["can_commit"] is False
    assert any(e["field"] == "last_name" for e in result["errors"])


def test_student_import_validation_accepts_sex_codes():
    headers, rows = parse_upload(io.BytesIO(b"first_name,last_name,sex\nJane,Doe,M\nJohn,Smith,F\n"))
    result = validate_rows(STUDENT_IMPORT_SPEC, headers, rows)
    assert result["valid_count"] == 2
    assert result["ready"][0]["gender"] == "male"
    assert result["ready"][1]["gender"] == "female"
    assert "email" not in result["ready"][0]


def test_student_import_rejects_invalid_sex():
    headers, rows = parse_upload(io.BytesIO(b"first_name,last_name,sex\nJane,Doe,X\n"))
    result = validate_rows(STUDENT_IMPORT_SPEC, headers, rows)
    assert result["valid_count"] == 0
    assert any(e["field"] == "gender" for e in result["errors"])


@pytest.mark.django_db
def test_student_import_requires_class_context(tenant, school_admin, academic_setup):
    _, school_class, _ = academic_setup
    row = {
        "first_name": "New",
        "last_name": "Student",
        "gender": "male",
    }
    context = StudentImportContext()
    errors = student_import_resolver(tenant, context=context, user=school_admin)(row, 2)
    assert any(e["field"] == "school_class" for e in errors)

    _, _, stream = academic_setup
    context = StudentImportContext(
        school_class_id=str(school_class.id),
        stream_id=str(stream.id),
    )
    errors = student_import_resolver(tenant, context=context, user=school_admin)(row, 2)
    assert errors == []
    assert row["_school_class_id"] == str(school_class.id)
    assert row["_stream_id"] == str(stream.id)


@pytest.mark.django_db
def test_student_import_requires_stream_when_class_has_streams(tenant, school_admin, academic_setup):
    _, school_class, _ = academic_setup
    row = {
        "first_name": "New",
        "last_name": "Student",
        "gender": "female",
    }
    context = StudentImportContext(school_class_id=str(school_class.id))
    errors = student_import_resolver(tenant, context=context, user=school_admin)(row, 2)
    assert any(e["field"] == "stream" for e in errors)


@pytest.mark.django_db
def test_commit_student_rows_assigns_class_and_stream(tenant, school_admin, academic_setup):
    _, school_class, stream = academic_setup
    rows = [{
        "first_name": "Peter",
        "last_name": "Kamau",
        "gender": "male",
        "_school_class_id": str(school_class.id),
        "_stream_id": str(stream.id),
    }]
    result = commit_student_rows(tenant, rows, actor=school_admin)
    assert result["created"] == 1
    student = Student.objects.get(tenant=tenant, first_name="Peter")
    assert student.school_class_id == school_class.id
    assert student.stream_id == stream.id
    assert student.gender == "male"
    assert student.date_of_birth == IMPORT_PLACEHOLDER_DOB
    assert is_profile_incomplete(student) is True


@pytest.mark.django_db
def test_commit_parent_rows(tenant, school_admin):
    rows = [{
        "first_name": "Jane",
        "last_name": "Wanjiku",
        "email": "jane.parent@test.edu",
        "phone": "0712345678",
    }]
    result = commit_parent_rows(tenant, rows, actor=school_admin)
    assert result["created"] == 1
    assert Parent.objects.filter(tenant=tenant, email="jane.parent@test.edu").exists()


@pytest.mark.django_db
def test_student_import_template_endpoint_defaults_to_excel(school_admin, import_plan):
    client = APIClient()
    client.force_authenticate(user=school_admin)
    response = client.get("/api/v1/students/import-template/")
    assert response.status_code == 200
    assert response["Content-Type"] == EXCEL_CONTENT_TYPE
    assert response.content[:2] == b"PK"
    assert "students_import_template.xlsx" in response["Content-Disposition"]


@pytest.mark.django_db
def test_student_validate_import_accepts_excel(school_admin, import_plan, academic_setup):
    _, school_class, stream = academic_setup
    client = APIClient()
    client.force_authenticate(user=school_admin)
    content, _, _ = generate_import_template(STUDENT_IMPORT_SPEC)
    upload = io.BytesIO(content)
    upload.name = "students.xlsx"
    response = client.post(
        "/api/v1/students/validate-import/",
        {
            "file": upload,
            "school_class": str(school_class.id),
            "stream": str(stream.id),
        },
        format="multipart",
    )
    assert response.status_code == 200
    data = response.json()["data"]
    assert data["valid_count"] == 0
    assert data["can_commit"] is False


@pytest.mark.django_db
def test_student_validate_import_endpoint(school_admin, import_plan, academic_setup):
    _, school_class, stream = academic_setup
    client = APIClient()
    client.force_authenticate(user=school_admin)
    csv_content = (
        "First Name,Last Name,Sex\n"
        "Peter,Kamau,M\n"
        "Mary,Wanjiku,F\n"
    )
    upload = io.BytesIO(csv_content.encode("utf-8"))
    upload.name = "students.csv"
    response = client.post(
        "/api/v1/students/validate-import/",
        {
            "file": upload,
            "school_class": str(school_class.id),
            "stream": str(stream.id),
        },
        format="multipart",
    )
    assert response.status_code == 200
    data = response.json()["data"]
    assert data["valid_count"] == 2
    assert data["can_commit"] is True
    assert data["import_context"]["school_class_name"] == "Grade 7"


@pytest.mark.django_db
def test_student_list_reports_incomplete_profiles(school_admin, import_plan, academic_setup):
    _, school_class, _ = academic_setup
    Student.objects.create(
        tenant=school_admin.tenant,
        admission_number="INC001",
        first_name="Incomplete",
        last_name="Student",
        date_of_birth=IMPORT_PLACEHOLDER_DOB,
        gender="male",
        enrollment_date="2026-01-15",
        school_class=school_class,
        created_by=school_admin,
        updated_by=school_admin,
    )
    client = APIClient()
    client.force_authenticate(user=school_admin)
    response = client.get("/api/v1/students/")
    assert response.status_code == 200
    payload = response.json()
    assert payload.get("meta", {}).get("incomplete_profile_count", 0) >= 1


@pytest.mark.django_db
def test_parent_import_template_endpoint(school_admin, import_plan):
    client = APIClient()
    client.force_authenticate(user=school_admin)
    response = client.get("/api/v1/students/parents/import-template/?file_format=xlsx")
    assert response.status_code == 200
    assert response["Content-Type"] == EXCEL_CONTENT_TYPE
    assert response.content[:2] == b"PK"
    assert "parents_import_template.xlsx" in response["Content-Disposition"]


@pytest.mark.django_db
def test_staff_import_template_endpoint_for_school_admin(school_admin, import_plan):
    client = APIClient()
    client.force_authenticate(user=school_admin)
    response = client.get("/api/v1/staff/import-template/")
    assert response.status_code == 200
    assert response["Content-Type"] == EXCEL_CONTENT_TYPE
    assert response.content[:2] == b"PK"
    assert "staff_import_template.xlsx" in response["Content-Disposition"]


@pytest.mark.django_db
def test_import_template_rejects_drf_format_query_param(school_admin, import_plan):
    """?format=xlsx collides with DRF format suffixes and must not be used."""
    client = APIClient()
    client.force_authenticate(user=school_admin)
    for path in ("/api/v1/staff/import-template/", "/api/v1/students/import-template/"):
        response = client.get(f"{path}?format=xlsx")
        assert response.status_code == 404


@pytest.mark.django_db
def test_staff_import_template_endpoint_for_hr_manager(db, tenant, import_plan, monkeypatch):
    class _EmailResult:
        success = True
        message = ""

    monkeypatch.setattr(
        "apps.staff.portal_credentials.EmailService.send",
        lambda *args, **kwargs: _EmailResult(),
    )
    hr_staff = onboard_staff(
        tenant,
        data={
            "first_name": "HR",
            "last_name": "Manager",
            "email": "hr.manager.bulk@test.edu",
            "phone": "+254700000701",
            "portal_role": UserRole.HR_MANAGER,
            "date_joined": "2026-01-01",
        },
    )
    client = APIClient()
    client.force_authenticate(user=hr_staff.user)
    response = client.get("/api/v1/staff/import-template/")
    assert response.status_code == 200, response.content[:300]
    assert response.content[:2] == b"PK"

    content, filename, _ = generate_import_template(STAFF_IMPORT_SPEC)
    assert filename == "staff_import_template.xlsx"
    assert len(content) > 0


@pytest.mark.django_db
def test_import_context_endpoint(school_admin, import_plan, academic_setup):
    client = APIClient()
    client.force_authenticate(user=school_admin)
    response = client.get("/api/v1/students/import-context/")
    assert response.status_code == 200
    classes = response.json()["data"]["classes"]
    assert len(classes) >= 1
    assert classes[0]["has_streams"] is True